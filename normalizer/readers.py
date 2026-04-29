"""Utilities for reading Excel workbooks."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import Cell


def load_workbook_with_formulas(path: str | Path) -> Workbook:
    """Load a workbook keeping formula strings intact (data_only=False)."""
    return openpyxl.load_workbook(str(path), data_only=False)


def get_sheet_names(wb: Workbook) -> list[str]:
    """Return the sheet names of a workbook."""
    return list(wb.sheetnames)


def get_headers(sheet: Any, header_row: int = 1) -> dict[int, str | None]:
    """Return a mapping of {column_index: header_value} for a given row."""
    headers: dict[int, str | None] = {}
    for cell in sheet[header_row]:
        value = cell.value
        if value is not None:
            headers[cell.column] = str(value).strip()
        else:
            headers[cell.column] = None
    return headers


def cell_address(row: int, col: int) -> str:
    """Convert (row, col) to an Excel-style address like A1."""
    return f"{openpyxl.utils.get_column_letter(col)}{row}"


def is_formula(cell: Cell) -> bool:
    """Return True if the cell contains a formula."""
    return isinstance(cell.value, str) and cell.value.startswith("=")


def normalize_formula(formula: str, current_col: int = 0, current_row: int = 0) -> str:
    """Normalise a formula for structural comparison.

    The following transformations are applied:

    * Case-folded to upper-case so comparisons are case-insensitive.
    * Single quotes around sheet names are stripped (Excel adds them when the
      name starts with a digit or contains special characters), so
      ``'1_Sheet'!A1`` and ``1_SHEET!A1`` compare as equal.
    * When *current_row* is given (> 0), same-sheet row numbers are converted
      to signed offsets relative to *current_row* (e.g. ``=A11`` in row 11
      becomes ``=A[+0]``, ``=A10`` in row 11 becomes ``=A[-1]``).  This lets
      the same structural formula copied across rows compare equal while still
      detecting genuine row-reference changes within a row.  Without
      *current_row*, row numbers are replaced with the flat placeholder ``#``.
    * When *current_col* is given (> 0), non-absolute same-sheet column
      references are converted to signed offsets relative to *current_col*
      (e.g. ``=D2`` in column C becomes ``=[+1][±row]``).  This lets equivalent
      formulas compare equal even when the column containing the formula
      differs between the template and the user workbook.  Cross-sheet
      references and absolute column references (``$D2``) are left as-is so
      intentional fixed anchors are still detected as mismatches.
    """
    if not formula:
        return formula
    # Upper-case for case-insensitive comparison
    formula = formula.upper()
    # Strip single quotes around sheet names (e.g. '1_Sheet'!A1 → 1_SHEET!A1)
    formula = re.sub(r"'([^']+)'!", r"\1!", formula)

    if current_col > 0:
        # ------------------------------------------------------------------ #
        # Step 1 – Normalise cross-sheet cell/range references first so their
        # column letters and row numbers are NOT relativised in steps 2-3.
        # Pattern matches:  SHEETNAME!<cell_or_range>
        # e.g.  1_SHEET!$B$38  or  SHEET!A1:B2
        # Cross-sheet rows reference absolute positions in another sheet and
        # must never be made relative to the current row.
        # ------------------------------------------------------------------ #
        def _normalize_xref(m: re.Match) -> str:
            # Replace row numbers inside the cross-sheet ref with # (absolute)
            return re.sub(r"(\$?[A-Z]+)\$?[0-9]+", r"\1#", m.group(0))

        formula = re.sub(
            r"[A-Z0-9_]+!\$?[A-Z]+\$?[0-9]+(?::\$?[A-Z]+\$?[0-9]+)?",
            _normalize_xref,
            formula,
        )

        # ------------------------------------------------------------------ #
        # Step 2 – Relativise non-absolute same-sheet column references and,
        # when current_row is known, the row number as well.
        # A relative column ref is NOT preceded by $ (absolute column) or !
        # (cross-sheet ref, already handled above).
        # ------------------------------------------------------------------ #
        def _relativize(m: re.Match) -> str:
            col_num = openpyxl.utils.column_index_from_string(m.group(1))
            col_offset = col_num - current_col
            if current_row > 0:
                row_offset = int(m.group(2)) - current_row
                return f"[{col_offset:+d}][{row_offset:+d}]"
            return f"[{col_offset:+d}]#"

        formula = re.sub(r"(?<![!$])([A-Z]+)\$?([0-9]+)", _relativize, formula)

        # Step 3 – Replace row numbers in any remaining absolute column refs
        # (e.g. $D$2) that survived steps 1-2.
        if current_row > 0:
            def _abs_col_row(m: re.Match) -> str:
                row_offset = int(m.group(2)) - current_row
                return f"{m.group(1)}[{row_offset:+d}]"
            formula = re.sub(r"(\$[A-Z]+)\$?([0-9]+)", _abs_col_row, formula)
        else:
            formula = re.sub(r"(\$[A-Z]+)\$?[0-9]+", r"\1#", formula)
    else:
        # Without column context: replace row numbers (relatively if current_row known)
        if current_row > 0:
            def _rel_row_only(m: re.Match) -> str:
                row_offset = int(m.group(2)) - current_row
                return f"{m.group(1)}[{row_offset:+d}]"
            formula = re.sub(r"(\$?[A-Z]+)\$?([0-9]+)", _rel_row_only, formula)
        else:
            formula = re.sub(r"(\$?[A-Z]+)\$?[0-9]+", r"\1#", formula)

    return formula
