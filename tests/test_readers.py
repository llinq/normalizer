"""Tests for normalizer.readers utilities."""

from __future__ import annotations

import io

import openpyxl
import pytest

from normalizer.readers import (
    get_headers,
    is_formula,
    normalize_formula,
)


class TestGetHeaders:
    def _make_sheet(self, row: list):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(row)
        return ws

    def test_returns_column_header_map(self):
        ws = self._make_sheet(["Name", "Value", "Tax"])
        headers = get_headers(ws, header_row=1)
        assert headers[1] == "Name"
        assert headers[2] == "Value"
        assert headers[3] == "Tax"

    def test_strips_whitespace(self):
        ws = self._make_sheet(["  Name  ", " Value"])
        headers = get_headers(ws, header_row=1)
        assert headers[1] == "Name"
        assert headers[2] == "Value"

    def test_none_values_included(self):
        ws = self._make_sheet(["Name", None, "Tax"])
        headers = get_headers(ws, header_row=1)
        assert headers[2] is None


class TestIsFormula:
    def _make_cell(self, value):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = value
        return ws["A1"]

    def test_formula_detected(self):
        cell = self._make_cell("=SUM(A1:A10)")
        assert is_formula(cell) is True

    def test_plain_string_not_formula(self):
        cell = self._make_cell("hello")
        assert is_formula(cell) is False

    def test_number_not_formula(self):
        cell = self._make_cell(42)
        assert is_formula(cell) is False

    def test_none_not_formula(self):
        cell = self._make_cell(None)
        assert is_formula(cell) is False


class TestNormalizeFormula:
    def test_row_numbers_replaced(self):
        assert normalize_formula("=SUM(B2:B100)") == "=SUM(B#:B#)"

    def test_absolute_refs_normalised(self):
        result = normalize_formula("=$B$2+$C$3")
        assert "#" in result
        assert "2" not in result
        assert "3" not in result

    def test_function_name_preserved(self):
        result = normalize_formula("=VLOOKUP(A2,B2:C10,2,0)")
        assert "VLOOKUP" in result

    def test_empty_formula_passthrough(self):
        assert normalize_formula("") == ""

    def test_case_insensitive(self):
        assert normalize_formula("=sum(b2:b10)") == normalize_formula("=SUM(B2:B10)")

    def test_relative_col_same_col(self):
        """Column letter at the same position as current_col has offset 0."""
        result = normalize_formula("=C2", current_col=3)  # C=3
        assert result == "=[+0]#"

    def test_relative_col_positive_offset(self):
        """Column D in current col C gives offset +1."""
        assert normalize_formula("=D2", current_col=3) == "=[+1]#"

    def test_relative_col_negative_offset(self):
        """Column A in current col C gives offset -2."""
        assert normalize_formula("=A2", current_col=3) == "=[-2]#"

    def test_absolute_col_not_relativized(self):
        """Absolute column refs ($D$2) are kept as-is (row replaced only)."""
        result = normalize_formula("=$D$2", current_col=3)
        assert result == "=$D#"
        assert "[" not in result

    def test_cross_sheet_col_not_relativized(self):
        """Cross-sheet column refs are not relativized."""
        result = normalize_formula("=Sheet1!$B$38", current_col=3)
        assert "[" not in result
        assert "$B#" in result

    def test_same_formula_different_col_position(self):
        """Structurally identical formulas at different column positions compare equal."""
        # =D2&" - "&E2 in col C (offsets +1, +2) vs =R2&" - "&S2 in col Q (offsets +1, +2)
        tmpl = normalize_formula('=D2&" - "&E2', current_col=3)   # col C=3
        user = normalize_formula('=R2&" - "&S2', current_col=17)  # col Q=17
        assert tmpl == user

    def test_different_formula_structure_still_detected(self):
        """Formulas with genuinely different column offsets are still flagged."""
        tmpl = normalize_formula("=D2+E2", current_col=3)   # offsets +1, +2
        user = normalize_formula("=D2+E2", current_col=17)  # offsets -13, -12 (col Q)
        assert tmpl != user

    def test_quoted_cross_sheet_col_not_relativized(self):
        """Single-quoted cross-sheet refs (stripped first) are not relativized."""
        result = normalize_formula("='1_Sheet'!$B$38", current_col=3)
        assert "[" not in result
        assert "$B#" in result
